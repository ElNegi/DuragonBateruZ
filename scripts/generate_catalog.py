from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import json,collections
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'outputs'/'CATALOGO_COMPLETO_CARTAS_NORMALIZADO.docx'; OUT.parent.mkdir(exist_ok=True)
cards=json.loads((ROOT/'src/cards/cards.json').read_text(encoding='utf8'))
d=Document(); sec=d.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
styles=d.styles; styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(9)
title=d.add_paragraph(style='Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER;title.add_run('Catalogo Completo de Cartas Normalizado')
d.add_paragraph('Ecos de Ki. Catálogo normalizado generado desde cards.json. Incluye las 370 cartas del documento fuente y preserva sus textos de efecto.')
for h,t in [('Sistema de dados y Ki','Cada ronda usa tantos D6 como el número de ronda, hasta seis. La suma individual de los dados genera Ki; el Ki gastado se resta del total disponible.'),('Tipos y etiquetas','Las cartas se organizan por tipo. Los ataques pueden tener Naturaleza Físico o Ki y Alcance Cercanía o Distancia. Las rarezas son Común, Poco común, Rara, Épica y Legendaria.'),('Palabras clave','Los campos de condiciones, objetivos, duración y palabras clave permiten al motor validar permanentes, reacciones, misiones, sagas, estados, raciales y combos.')]:d.add_heading(h,1);d.add_paragraph(t)
d.add_heading('Cantidades por tipo',1);table=d.add_table(rows=1,cols=2);table.style='Table Grid';table.rows[0].cells[0].text='Tipo';table.rows[0].cells[1].text='Cantidad';counts=collections.Counter(c['sourceType'] for c in cards)
for k,v in counts.items():r=table.add_row().cells;r[0].text=k;r[1].text=str(v)
for typ in counts:
 d.add_page_break();d.add_heading(typ,1)
 for c in [x for x in cards if x['sourceType']==typ]:
  d.add_heading(c['name'],2);p=d.add_paragraph();p.add_run('Tipo: ').bold=True;p.add_run(c['cardType']);p.add_run('   Coste: ').bold=True;p.add_run(str(c['cost']));p.add_run('   Rareza: ').bold=True;p.add_run(c['rarity'])
  fields=[('Raza',c.get('raceRestriction','No aplica')),('Naturaleza',c.get('nature','No aplica')),('Alcance',c.get('range','No aplica')),('Daño',c.get('damage','No aplica')),('Efecto',c['rulesText']),('Condición',c['conditions'] or 'No aplica'),('Duración',c['duration']),('Objetivos',c['targeting']),('Notas',', '.join(c['keywords']))]
  for label,value in fields:p=d.add_paragraph(style='Normal');p.paragraph_format.left_indent=Inches(.2);p.add_run(label+': ').bold=True;p.add_run(str(value))
d.save(OUT);print(OUT)
